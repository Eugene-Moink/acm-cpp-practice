# -*- coding: utf-8 -*-
"""
按《计算机学院实验报告模板》生成《算法设计与分析》实验一实验报告。
模板来源：计算机学院实验报告模板－数据库系统概论.docx（课程名称改为“算法设计与分析”）
内容来源：干林锋-蛮力算法的设计与性能分析-实验报告.docx（正文、数据、18 张截图）
"""
import os
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TPL = os.path.join(ROOT, "计算机学院实验报告模板－数据库系统概论.docx")
OUT = os.path.join(ROOT, "干林锋-实验一-蛮力算法的设计与性能分析-实验报告.docx")
IMGDIR = os.path.join(ROOT, "img")

SONG = "宋体"
HEI = "黑体"
KAI = "楷体"
EN = "Times New Roman"
MONO = "Consolas"
BODY = 10.5          # 五号
SMALL = 9.5
CAPTION = 9.0

CENTER = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT
JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY


# ----------------------------------------------------------------------------- 基础工具
def set_run_font(run, cn=SONG, en=EN, size=BODY, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)


def add_par(container, text="", size=BODY, bold=False, cn=SONG, en=EN, align=None,
            indent_chars=0, before=0, after=0, line=1.25, keep_next=False, italic=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    if text:
        r = p.add_run(text)
        set_run_font(r, cn, en, size, bold, italic)
    return p


def h1(cell, text):
    return add_par(cell, text, size=12, bold=True, cn=HEI, align=LEFT,
                   before=8, after=4, line=1.25, keep_next=True)


def h2(cell, text):
    return add_par(cell, text, size=11, bold=True, cn=HEI, align=LEFT,
                   before=6, after=3, line=1.25, keep_next=True)


def h3(cell, text):
    return add_par(cell, text, size=BODY, bold=True, cn=HEI, align=LEFT,
                   before=4, after=2, line=1.25, keep_next=True)


def body(cell, text):
    return add_par(cell, text, indent_chars=2, align=JUSTIFY)


def plain(cell, text, size=BODY, indent=0):
    return add_par(cell, text, size=size, indent_chars=indent, align=JUSTIFY)


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag == qn("w:p"):
            tc.remove(child)


def del_par(par):
    par._element.getparent().remove(par._element)


def set_table_borders(table, sz=4):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "auto")
        borders.append(e)
    tblPr.append(borders)


TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
]


def normalize_tblPr(table):
    """按 OOXML 架构规定的顺序重排 w:tblPr 子元素，避免 Word 报文件损坏"""
    tblPr = table._tbl.tblPr
    children = list(tblPr)
    for el in children:
        tblPr.remove(el)

    def key(el):
        tag = el.tag.split("}")[-1]
        return TBLPR_ORDER.index(tag) if tag in TBLPR_ORDER else len(TBLPR_ORDER)

    for el in sorted(children, key=key):
        tblPr.append(el)


def set_cell_valign(cell, val="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:vAlign")):
        tcPr.remove(old)
    e = OxmlElement("w:vAlign")
    e.set(qn("w:val"), val)
    tcPr.append(e)


def fill_cell(cell, lines, size=BODY, bold=False, cn=SONG, en=EN, align=CENTER, line=1.0):
    """把若干行文本写入单元格（第一段复用，其余追加）"""
    if isinstance(lines, str):
        lines = [lines]
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    for i, txt in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        pf.line_spacing = line
        p.alignment = align
        if txt:
            r = p.add_run(txt)
            set_run_font(r, cn, en, size, bold)


def add_data_table(cell, headers, rows, widths, caption=None, size=SMALL,
                   header_size=None, aligns=None, caption_size=BODY):
    """在单元格中插入一张带边框的数据表（表题在表上方）"""
    if caption:
        add_par(cell, caption, size=caption_size, bold=True, cn=HEI, align=CENTER,
                before=8, after=2, line=1.0, keep_next=True)
    t = cell.add_table(rows=len(rows) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)

    total = sum(widths)
    tblPr = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(total * 567)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for r in t.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = Cm(w)

    hsize = header_size or size
    for i, htext in enumerate(headers):
        c = t.cell(0, i)
        set_cell_valign(c, "center")
        fill_cell(c, htext, size=hsize, bold=True, cn=HEI, align=CENTER)

    if aligns is None:
        aligns = [CENTER] * len(headers)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            set_cell_valign(c, "center")
            fill_cell(c, str(val), size=size, align=aligns[ci])

    # 表后留一点空隙
    add_par(cell, "", size=6, line=1.0, after=2)
    normalize_tblPr(t)
    return t


def add_figure(cell, imgname, caption, width_cm):
    p = cell.add_paragraph()
    p.alignment = CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    p.add_run().add_picture(os.path.join(IMGDIR, imgname), width=Cm(width_cm))
    cp = cell.add_paragraph()
    cp.alignment = CENTER
    cp.paragraph_format.space_after = Pt(8)
    cp.paragraph_format.line_spacing = 1.0
    r = cp.add_run(caption)
    set_run_font(r, SONG, EN, CAPTION)


def add_code(cell, code_text, size=8.5):
    for ln in code_text.split("\n"):
        p = cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.first_line_indent = Pt(0)
        r = p.add_run(ln)
        set_run_font(r, SONG, MONO, size)
    add_par(cell, "", size=6, line=1.0)


def read_text(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read().rstrip("\n")


def replace_run_text(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    return False


def fill_blank_run(par, value, pad=20):
    """把段落中最长的一段纯空白 run 填成给定内容"""
    best = None
    for r in par.runs:
        if r.text.strip() == "" and (best is None or len(r.text) > len(best.text)):
            best = r
    if best is None:
        return False
    tail = max(2, pad - len(value) * 2)
    best.text = " " + value + " " * tail
    return True


# ----------------------------------------------------------------------------- 打开模板
doc = docx.Document(TPL)

# 1) 封面：课程名称改为算法设计与分析，填写班级/学号/姓名/指导教师
#    做法：保留原段落第一个 run 的格式，重写为“标签 + 对齐空格 + 值”，使各行取值列对齐
import copy as _copy

COVER_LABELS = [
    ("课程名称", "课程名称："),
    ("班", "班    级："),
    ("学", "学    号："),
    ("姓", "姓    名："),
    ("指导教师", "指导教师："),
]
COVER_VALUES = {
    "课程名称": "算法设计与分析",
    "班": "计科2班",
    "学": "25080900217",
    "姓": "干林锋",
    "指导教师": "郭顺超",
}


def match_cover_key(txt):
    if txt.startswith("课程名称"):
        return "课程名称"
    if txt.startswith("班") and "级：" in txt:
        return "班"
    if txt.startswith("学") and "号：" in txt:
        return "学"
    if txt.startswith("姓") and "名：" in txt:
        return "姓"
    if txt.startswith("指导教师"):
        return "指导教师"
    return None


for p in doc.paragraphs:
    key = match_cover_key(p.text)
    if key is None or not p.runs:
        continue
    label = dict(COVER_LABELS)[key]
    value = COVER_VALUES[key]
    first = p.runs[0]
    rPr = first._element.find(qn("w:rPr"))
    for r in list(p.runs)[1:]:
        r._element.getparent().remove(r._element)
    first.text = label
    pad = max(2, 12 - len(label))
    nr = p.add_run(" " * pad + value)
    if rPr is not None:
        nr._element.insert(0, _copy.deepcopy(rPr))

table = doc.tables[0]

# 2) 表头信息
fill_cell(table.rows[0].cells[2], "蛮力算法的设计与性能分析", size=10.5, align=CENTER)
fill_cell(table.rows[0].cells[6], "4", size=10.5, align=CENTER)
fill_cell(table.rows[1].cells[1], "2026.09.21", size=10.5, align=CENTER)
# 实验地点、实验成绩留空由教师填写

# ----------------------------------------------------------------------------- 实验目的
c = table.rows[3].cells[1]
clear_cell(c)
h3(c, "一、实验题目1：冒泡、选择、归并排序算法的设计与分析")
for t in [
    "1. 深入理解蛮力法思想在排序问题中的应用，掌握冒泡排序、选择排序的核心逻辑与执行流程，明确其时间复杂度特性。",
    "2. 理解分治法思想在归并排序中的体现，掌握“分解—合并”流程及有序子数组的合并方法，对比蛮力排序与分治排序的效率差异。",
    "3. 实现随机数据生成模块，生成不同规模（100 个元素、1000 个元素、10000 个元素）的随机整数数组。",
    "4. 通过上述三种规模的随机整数数组进行测试，统计三种排序算法的执行时间、比较次数与移动次数，分析算法性能随数据规模变化的规律，提升算法性能分析与 C/C++ 编程实践能力。",
]:
    plain(c, t)
h3(c, "二、实验题目2：蛮力法与改进版百鸡问题算法设计与分析")
for t in [
    "1. 理解蛮力法“枚举所有可能解 + 验证约束条件”的核心思想，掌握问题边界分析与枚举范围优化方法。",
    "2. 对比基础蛮力法与改进蛮力法的时间复杂度差异，培养“通过约束条件减少枚举量”的算法优化思维。",
    "3. 熟练运用 C/C++ 的循环结构、数组存储、输入输出操作，实现多解问题的求解与结果展示。",
    "4. 学会设计测试用例验证算法正确性，通过运行时间对比分析优化算法的实际效果。",
]:
    plain(c, t)

# ----------------------------------------------------------------------------- 实验内容
c = table.rows[4].cells[1]
clear_cell(c)

h1(c, "一、实验题目1：冒泡、选择、归并排序算法的设计与分析")

h2(c, "（一）算法原理梳理")

h3(c, "1. 冒泡排序（改进版）")
body(c, "基本原理：冒泡排序属于典型的蛮力法。它反复扫描待排序序列，依次比较相邻的两个元素，若前者大于后者（逆序）则交换两者位置。这样每完成一趟扫描，当前待排序区间中的最大元素就会像水中的气泡一样“浮”到区间末尾，有序区长度增加 1。若序列长度为 n，则最多需要 n−1 趟扫描，每趟扫描的比较次数依次为 n−1、n−2、…、1。")
body(c, "改进思路：若某一趟扫描中“没有发生任何元素交换”，说明整个序列已经有序，后续趟次不可能再产生交换，因此立即终止算法。这一改进使算法在处理基本有序或已然有序的序列时性能显著提升。")
body(c, "实测验证：以本实验 n = 100 的随机数组为例，数据本身无序，改进版共执行 4 797 次比较，小于最坏情况的 4 950 次，说明确实有若干趟扫描在序列已有序时提前终止，改进生效。")
add_data_table(
    c,
    ["分析项", "结论"],
    [
        ["时间复杂度（最坏、平均）", "O(n²)，每趟比较次数呈等差数列，总计 n(n−1)/2 次"],
        ["时间复杂度（最好）", "O(n)，原始序列已有序时，第 1 趟无交换即终止"],
        ["空间复杂度", "O(1)，仅需常数个辅助变量，原地排序"],
        ["稳定性", "稳定，仅当 a[j] > a[j+1] 时才交换，相等元素相对次序不变"],
        ["统计量", "比较次数约 n²/2；移动次数 = 3 × 交换次数"],
    ],
    widths=[3.6, 9.8],
    caption="表 1-1  冒泡排序（改进版）性能特征分析",
    aligns=[CENTER, LEFT],
)

h3(c, "2. 选择排序")
body(c, "基本原理：选择排序把数组划分为“已排序区”和“未排序区”，初始时已排序区为空。每一趟从未排序区中选出最小（或最大）的元素，将其与未排序区的第一个元素交换位置，使已排序区长度增加 1，n−1 趟之后整个序列有序。")
body(c, "特点：选择排序的比较次数与数据的初始排列无关，第 i 趟固定比较 n−1−i 次，总计恒为 n(n−1)/2 次；但每趟最多只发生 1 次交换，因此元素移动次数很少，仅为 O(n) 量级。这一点在本实验数据中体现得非常明显：n = 10000 时选择排序的比较次数为 49 995 000 次，移动次数却只有 29 979 次。")
add_data_table(
    c,
    ["分析项", "结论"],
    [
        ["时间复杂度（最好、最坏、平均）", "均为 O(n²)，比较次数恒为 n(n−1)/2"],
        ["空间复杂度", "O(1)，原地排序"],
        ["稳定性", "不稳定，交换可能破坏相等元素的相对次序"],
        ["统计量特点", "比较次数固定；移动次数 = 3 × 交换次数，仅为 O(n)"],
    ],
    widths=[4.6, 8.8],
    caption="表 1-2  选择排序性能特征分析",
    aligns=[CENTER, LEFT],
)

h3(c, "3. 归并排序")
body(c, "基本原理：归并排序是分治法的典型应用，其执行过程可概括为“分解—合并”两个阶段。")
body(c, "分解：把待排序区间从中间一分为二（取 mid = left + (right−left)/2），对左右两个子区间分别递归地排序；当子区间只剩 1 个元素时天然有序，递归结束。")
body(c, "合并：把两个已经有序的子数组 a[left..mid] 与 a[mid+1..right] 合并为一个有序数组。做法是设置两个游标分别指向两段的起始位置，每次比较两个游标所指元素，把较小者放入临时数组，游标后移；当某一段取完后，把另一段剩余元素依次搬入临时数组；最后把临时数组的结果整体写回原数组对应区间。合并过程中若取 a[i] <= a[j] 则优先取左段元素，从而保证算法的稳定性。")
body(c, "复杂度分析：分解过程把区间二分，共 log₂n 层；每层的合并操作需要遍历该层全部 n 个元素，代价为 O(n)。因此无论数据初始排列如何，总时间复杂度恒为 O(n log n)，这是归并排序相对两种蛮力排序的根本优势。代价是需要一个与原数组等长的辅助数组，空间复杂度为 O(n)。")
add_data_table(
    c,
    ["分析项", "结论"],
    [
        ["时间复杂度（最好、最坏、平均）", "均为 O(n log n)，与数据初始排列无关"],
        ["空间复杂度", "O(n)，需要等长辅助数组存放合并结果"],
        ["稳定性", "稳定，合并时取 a[i] <= a[j] 保证相等元素先后次序不变"],
        ["统计量", "比较次数约 n·log₂n；移动次数约 2n·log₂n（含写回原数组）"],
    ],
    widths=[4.6, 8.8],
    caption="表 1-3  归并排序性能特征分析",
    aligns=[CENTER, LEFT],
)

h2(c, "（二）测试用例设计")
body(c, "为验证三种算法在不同数据规模下的性能差异，设计如下测试用例。三种规模使用同一随机数生成规则（srand/rand 生成 [1, 100000] 的整数），每种算法独立执行 3 次，取 3 次的平均值作为最终性能数据。")
add_data_table(
    c,
    ["数据规模", "测试次数", "三种算法的比较结论（运行时间上的差异）"],
    [
        ["100 个元素", "3 次（取平均）", "归并排序最快（1.33 微秒）；冒泡、选择排序分别是它的 5.0 倍和 6.0 倍，规模小时差距不大"],
        ["1000 个元素", "3 次（取平均）", "归并排序（26.66 微秒）已明显领先，比冒泡排序快约 30.4 倍、比选择排序快约 28.4 倍"],
        ["10000 个元素", "3 次（取平均）", "归并排序（333.43 微秒）比冒泡排序快约 186.2 倍、比选择排序快约 276.5 倍，优势急剧扩大"],
    ],
    widths=[2.2, 2.2, 9.0],
    caption="表 1-4  三种排序算法的测试用例设计",
    aligns=[CENTER, CENTER, LEFT],
)
plain(c, "补充说明：随机数使用固定种子 srand(20250809 + n) 生成，故同一规模下每次运行得到的数组完全相同，实验结果可复现；由于三种算法面对的数据完全一致，比较次数、移动次数与耗时的对比是公平有效的。")

h2(c, "（三）运行截图")
body(c, "下列截图取自程序在 n = 100、1000、10000 三种规模下的真实运行输出，每张截图均显示了数据规模、算法名称、排序前后数组片段、3 次运行的比较次数、移动次数与平均执行时间，以及结果正确性校验。")
plain(c, "（1）数据规模 n = 100", indent=0)
add_figure(c, "exp1_n100_1.png", "图 1-1  冒泡排序（改进版）在 n = 100 时的运行结果", 13.4)
add_figure(c, "exp1_n100_2.png", "图 1-2  选择排序在 n = 100 时的运行结果", 13.4)
add_figure(c, "exp1_n100_3.png", "图 1-3  归并排序在 n = 100 时的运行结果", 13.4)
plain(c, "（2）数据规模 n = 1000", indent=0)
add_figure(c, "exp1_n1000_1.png", "图 1-4  冒泡排序（改进版）在 n = 1000 时的运行结果", 13.4)
add_figure(c, "exp1_n1000_2.png", "图 1-5  选择排序在 n = 1000 时的运行结果", 13.4)
add_figure(c, "exp1_n1000_3.png", "图 1-6  归并排序在 n = 1000 时的运行结果", 13.4)
plain(c, "（3）数据规模 n = 10000", indent=0)
add_figure(c, "exp1_n10000_1.png", "图 1-7  冒泡排序（改进版）在 n = 10000 时的运行结果", 13.4)
add_figure(c, "exp1_n10000_2.png", "图 1-8  选择排序在 n = 10000 时的运行结果", 13.4)
add_figure(c, "exp1_n10000_3.png", "图 1-9  归并排序在 n = 10000 时的运行结果", 13.4)

h2(c, "（四）性能数据汇总")
body(c, "将 3 次运行的平均值汇总如下表。执行时间单位为微秒（μs），比较次数与移动次数为单次排序的统计值。")
add_data_table(
    c,
    ["数据规模 n", "算法", "比较次数（次）", "移动次数（次）", "平均执行时间（μs）"],
    [
        ["100", "冒泡排序（改进版）", "4,797", "7,119", "6.666"],
        ["100", "选择排序", "4,950", "291", "8.000"],
        ["100", "归并排序", "536", "1,344", "1.333"],
        ["1000", "冒泡排序（改进版）", "499,004", "742,056", "810.229"],
        ["1000", "选择排序", "499,500", "2,982", "756.873"],
        ["1000", "归并排序", "8,691", "19,952", "26.661"],
        ["10000", "冒泡排序（改进版）", "49,984,989", "74,503,941", "62,073.167"],
        ["10000", "选择排序", "49,995,000", "29,979", "92,195.267"],
        ["10000", "归并排序", "120,493", "267,232", "333.433"],
    ],
    widths=[1.9, 3.3, 2.8, 2.8, 2.6],
    caption="表 1-5  三种排序算法在三种数据规模下的平均性能数据",
    size=9.0,
    header_size=9.0,
)

h2(c, "（五）性能对比图表")
add_figure(c, "chart_time.png", "图 1-10  三种排序算法执行时间随数据规模的变化（双对数坐标）", 13.4)
add_figure(c, "chart_counts.png", "图 1-11  三种排序算法的比较次数与移动次数对比（对数坐标）", 13.4)

h2(c, "（六）性能随数据规模变化的规律分析")
for t in [
    "（1）三种算法耗时随规模的增长速度完全不同。规模由 100 增至 10000（扩大 100 倍）时，冒泡排序耗时由 6.67 微秒增至 62 073 微秒，选择排序由 8.00 微秒增至 92 195 微秒，增长幅度都在 7 000 倍以上，符合 O(n²) 的平方级增长规律；而归并排序由 1.33 微秒增至 333 微秒，只增长约 250 倍，与 O(n log n) 的理论预期一致。",
    "（2）数据规模越大，归并排序的优势越明显。n = 100 时归并排序（1.33 微秒）只是略快于冒泡排序（6.67 微秒），三种算法处于同一数量级；n = 1000 时归并排序已比冒泡排序快约 30.4 倍；到 n = 10000 时，归并排序（333 微秒）比冒泡排序快约 186.2 倍、比选择排序快约 276.5 倍。这说明规模较小时三种算法差距不明显，规模增大后 O(n log n) 与 O(n²) 之间的差距才充分体现出来。",
    "（3）比较次数与移动次数由算法设计决定，与运行环境无关。n = 10000 时，冒泡排序比较 49 984 989 次、选择排序比较 49 995 000 次，都接近 n(n−1)/2 = 49 995 000 次的理论值，而归并排序只需 120 493 次；移动次数方面选择排序最少（29 979 次，每趟最多交换一次），冒泡排序最多（74 503 941 次），归并排序居中（267 232 次）。",
    "（4）改进版冒泡排序的提前终止确实有效，但对随机数据收益有限。n = 100 时最坏需比较 4 950 次，实际为 4 797 次；n = 1000 时最坏 499 500 次，实际 499 004 次；n = 10000 时最坏 49 995 000 次，实际 49 984 989 次，均少于最坏情况，说明有若干趟扫描提前结束。相对最坏情况的减少量分别为 3.1%、0.10% 和 0.02%，呈下降趋势——随机数据整体无序，规模越大越难出现“整趟无交换”的情形。",
]:
    plain(c, t)

h2(c, "（七）实验结论")
for t in [
    "（1）三种算法都能正确完成排序，程序内置的“非递减有序”校验在全部 9 组测试中全部通过，说明算法实现正确。",
    "（2）冒泡排序与选择排序的时间复杂度为 O(n²)，n = 10000 时耗时分别约为 62 073 微秒和 92 195 微秒；归并排序的分治策略把复杂度降到 O(n log n)，耗时仅 333 微秒，比冒泡排序快约 186.2 倍、比选择排序快约 276.5 倍，且规模越大优势越明显。",
    "（3）比较次数与移动次数由算法设计决定：选择排序比较次数固定但移动次数最少，冒泡排序移动次数最多，归并排序两项指标都远小于两种蛮力排序。",
    "（4）改进版冒泡排序的提前终止机制在近似有序的数据上能明显减少比较次数，但对完全随机的大规模数据改善有限。",
    "（5）实验表明，通过改进算法设计降低时间复杂度，是提升程序性能最有效的手段。",
]:
    plain(c, t)

# ------------------------------------------------------------------ 题目2
h1(c, "二、实验题目2：蛮力法与改进版百鸡问题算法设计与分析")

h2(c, "（一）问题描述与约束条件")
body(c, "问题描述：公鸡 5 钱 1 只，母鸡 3 钱 1 只，小鸡 1 钱 3 只。用 n 钱买 n 只鸡，问公鸡 a、母鸡 b、小鸡 c 各多少只？约束条件为：")
plain(c, "a ≥ 0，b ≥ 0，c ≥ 0，c % 3 == 0，a + b + c = n，5a + 3b + c/3 = n")
plain(c, "说明：本实验中“总钱数”与“总鸡数”取值相同，均为同一个参数 n，即所谓“用 n 钱买 n 只鸡”。")

h2(c, "（二）基础蛮力法")
body(c, "设计思路：直接枚举公鸡数 a 和母鸡数 b 的所有可能取值，利用 a + b + c = n 直接推出小鸡数 c = n − a − b，再验证约束条件是否全部满足。由于 a 与 b 的枚举范围都是 [0, n]，各自有 n+1 种取值，因此循环体恰好执行 (n+1)² = n² + 2n + 1 次，与具体数据无关。")
body(c, "约束验证：c ≥ 0；c % 3 == 0（小鸡必须 3 只一组）；5a + 3b + c/3 == n。为避免整数除法带来的误差，程序中把该式等价改写为 15a + 9b + c == 3n 进行判断。")
body(c, "复杂度：时间复杂度 O(n²)，空间复杂度 O(k)（k 为解的组数，用于存储解数组）。")

h2(c, "（三）改进蛮力法")
body(c, "改进的核心是“用约束条件化简枚举维度”，通过代数消元把一个变量用另一个变量表示出来，从而把双重循环降为单重循环。推导过程如下：")
plain(c, "由  a + b + c = n          ……（1）")
plain(c, "    5a + 3b + c/3 = n      ……（2）")
plain(c, "把 (1) 式乘以 3：3a + 3b + 3c = 3n      ……（3）")
plain(c, "把 (2) 式乘以 3：15a + 9b + c = 3n      ……（4）")
plain(c, "(4) − (3) 得：12a + 6b − 2c = 0，即 c = 6a + 3b      ……（5）")
plain(c, "把 (5) 代入 (1)：a + b + 6a + 3b = n，即 7a + 4b = n，于是 b = (n − 7a) / 4      ……（6）")
plain(c, "再由 c = n − a − b 即可算出小鸡数。可见只需枚举一个变量 a，b 与 c 都能直接算出。")
body(c, "枚举范围优化：由 b = (n − 7a)/4 ≥ 0 可得 a ≤ n/7，因此 a 只需在 [0, n/7] 内枚举，循环体执行次数由 (n+1)² 次降到约 n/7 + 1 次。")
body(c, "复杂度：时间复杂度由 O(n²) 降为 O(n)，降低了一个数量级；空间复杂度仍为 O(k)。")

h2(c, "（四）测试用例设计")
body(c, "按实验要求设计 3 组测试用例，覆盖小规模、经典规模与大规模三种情形，用同一份程序分别运行验证，结果如下表。")
add_data_table(
    c,
    ["测试用例", "输入 n", "解的组数", "基础版循环次数", "改进版循环次数"],
    [
        ["1", "30", "1", "961", "5"],
        ["2", "100", "4（与预期一致）", "10,201", "15"],
        ["3", "1000", "36", "1,002,001", "143"],
    ],
    widths=[1.9, 1.9, 3.1, 3.3, 3.3],
    caption="表 2-1  百鸡问题测试用例与循环次数统计",
    aligns=[CENTER, CENTER, CENTER, CENTER, CENTER],
)
plain(c, "注：任务书中预估 n = 30 时有 2 组解，由约束 7a + 4b = 30 可知其非负整数解只有 a = 2、b = 4、c = 24 这一组，程序运行结果与理论推导完全一致，故实际为 1 组解。n = 100 时求得 4 组解：(0, 25, 75)、(4, 18, 78)、(8, 11, 81)、(12, 4, 84)，与经典百鸡问题的预期结果完全一致；改进蛮力法在三种规模下求得的解集与基础蛮力法完全相同，验证了约束化简推导的正确性。")

h2(c, "（五）运行截图")
body(c, "下列截图为程序在三种测试用例下的真实运行输出，均显示了输入 n、算法名称、求得的解、循环体执行次数与执行时间。")
plain(c, "（1）测试用例 n = 30", indent=0)
add_figure(c, "exp2_n30_basic.png", "图 2-1  基础蛮力法求解 n = 30 的结果", 13.0)
add_figure(c, "exp2_n30_impr.png", "图 2-2  改进蛮力法求解 n = 30 的结果", 13.0)
plain(c, "（2）测试用例 n = 100", indent=0)
add_figure(c, "exp2_n100_basic.png", "图 2-3  基础蛮力法求解 n = 100 的结果", 13.0)
add_figure(c, "exp2_n100_impr.png", "图 2-4  改进蛮力法求解 n = 100 的结果", 13.0)
plain(c, "（3）测试用例 n = 1000", indent=0)
add_figure(c, "exp2_n1000_basic.png", "图 2-5  基础蛮力法求解 n = 1000 的结果", 13.0)
add_figure(c, "exp2_n1000_impr.png", "图 2-6  改进蛮力法求解 n = 1000 的结果", 13.0)

h2(c, "（六）优化效果分析")
body(c, "两种算法在三种测试用例下的循环次数与执行时间对比如下表。")
add_data_table(
    c,
    ["n", "解数", "基础版循环次数", "改进版循环次数", "循环次数降低", "基础版耗时(μs)", "改进版耗时(μs)", "耗时降低"],
    [
        ["30", "1", "961", "5", "192.2 倍", "0.788", "0.003", "262.7 倍"],
        ["100", "4", "10,201", "15", "680.1 倍", "6.756", "0.008", "844.3 倍"],
        ["1000", "36", "1,002,001", "143", "7007.0 倍", "550.015", "0.101", "5471.1 倍"],
    ],
    widths=[0.9, 1.0, 2.0, 2.0, 1.9, 1.9, 1.9, 1.8],
    caption="表 2-2  基础蛮力法与改进蛮力法性能对比",
    size=8.5,
    header_size=8.5,
)
add_figure(c, "chart_exp2.png", "图 2-7  两种算法循环体执行次数与执行时间对比（对数坐标）", 13.4)
for t in [
    "（1）循环次数大幅减少。n = 30 时循环次数由 961 次降到 5 次，n = 100 时由 10 201 次降到 15 次，n = 1000 时由 1 002 001 次降到 143 次，分别降低 192.2 倍、680.1 倍和 7007.0 倍。这与理论分析一致：基础版循环 (n+1)² 次，改进版循环 n/7 + 1 次，规模越大降低越多。",
    "（2）执行时间也明显下降。n = 1000 时基础蛮力法单次耗时 550.015 微秒，改进蛮力法只需 0.101 微秒，降低约 5471.1 倍；n = 30 时两者耗时分别为 0.788 微秒和 0.003 微秒。规模越大，改进的效果越突出。",
    "（3）改进算法结果正确。三种测试用例下，改进蛮力法与基础蛮力法求得的解组数、每组解的数值完全一致（程序中逐组比对校验），说明消元推导和枚举范围 a ≤ n/7 都是正确的，没有遗漏合法解。",
]:
    plain(c, t)

h2(c, "（七）实验结论")
for t in [
    "（1）基础蛮力法“枚举所有组合 + 验证约束”的思路直观可靠，但双重循环使时间复杂度为 O(n²)，n = 1000 时循环次数高达 1 002 001 次。",
    "（2）改进蛮力法通过消元把三个未知量化为一个枚举变量，并把枚举范围压缩到 [0, n/7]，时间复杂度降为 O(n)，循环次数降到 n/7 + 1 量级。",
    "（3）实测数据显示，n = 1000 时改进算法循环次数降低 7007 倍、执行时间降低 5471 倍，且两种算法解集完全一致，说明“利用约束条件减少枚举量”能在保证正确性的同时大幅提升效率。",
    "（4）优化效果随规模增大而增强，说明在枚举类问题中，充分利用约束条件缩小搜索空间是最直接有效的优化途径。",
]:
    plain(c, t)

# ----------------------------------------------------------------------------- 实验设备及软件环境
c = table.rows[5].cells[1]
clear_cell(c)
h3(c, "（一）硬件环境")
for t in [
    "处理器：Intel(R) Core(TM) Ultra 9 275HX（24 个逻辑处理器）；体系结构：x86-64（AMD64）；内存：16 GB。",
]:
    plain(c, t)
h3(c, "（二）软件环境")
add_data_table(
    c,
    ["项目", "配置"],
    [
        ["操作系统", "Windows 11 家庭中文版（64 位，版本 25H2，内部版本 26200）"],
        ["编程语言", "C++17"],
        ["编译环境", "g++ (TDM64 MinGW-w64) 10.3.0"],
        ["编译选项", "-std=c++17 -O2"],
        ["计时方式", "C++11 <chrono> 高精度时钟（微秒级）"],
        ["代码编辑器", "Visual Studio Code"],
        ["数据整理与绘图", "Excel 表格整理 + Python 3.11 / Matplotlib 绘制双对数、对数坐标对比图"],
        ["截图工具", "Windows 截图工具（Snipping Tool）"],
    ],
    widths=[3.4, 10.0],
    caption=None,
    aligns=[CENTER, LEFT],
)

# ----------------------------------------------------------------------------- 实验方法及基本操作步骤
c = table.rows[7].cells[0]
clear_cell(c)
h2(c, "（一）总体实验流程")
steps = [
    ("步骤 1：环境准备。", "安装 TDM64 MinGW-w64（g++ 10.3.0）与 Visual Studio Code，确认 g++ 命令可用；创建实验目录，分别建立排序算法测试程序与百鸡问题求解程序两个源文件。"),
    ("步骤 2：随机数据生成（题目1）。", "编写 genRandomArray() 函数，用 srand(seed)/rand() 生成 [1, 100000] 的随机整数；种子固定为 20250809 + n，保证同一规模下每次运行得到完全相同的数组，实验结果可复现；三种算法共用同一份数据，保证比较公平。"),
    ("步骤 3：算法实现与统计埋点。", "分别实现改进版冒泡排序、选择排序和递归归并排序；统一“移动次数”口径：每一次对数组位置的写操作记 1 次（赋值记 1 次，一次交换记 3 次），归并排序把辅助数组写回原数组同样计入；比较次数在每个比较语句处累加。"),
    ("步骤 4：计时方法。", "小规模数据单次排序耗时常常不足 1 微秒，直接测量会被时钟分辨率淹没，因此对同一随机数组连续排序 repeat 次（n = 100、1000、10000 分别取 500、50、1 次），用总耗时除以 repeat 得到单次平均耗时，每种算法测量 3 次后取平均。"),
    ("步骤 5：正确性校验。", "编写 isSorted() 校验排序结果是否为非递减序列；百鸡问题中逐组比对基础蛮力法与改进蛮力法求得的解集是否完全一致，确保优化没有遗漏解。"),
    ("步骤 6：编译与运行。", "使用命令 g++ -std=c++17 -O2 源文件.cpp -o 目标程序.exe 编译；程序开头调用 SetConsoleOutputCP(CP_UTF8) 把控制台输出切换为 UTF-8，避免中文乱码；运行程序并按要求截图保存。"),
    ("步骤 7：测试用例输入（题目2）。", "分别输入 n = 30、100、1000 三组测试用例，记录两种算法求得的解的组数、每组解、循环体执行次数与单次执行时间。"),
    ("步骤 8：数据整理与分析。", "把 3 次运行的比较次数、移动次数与平均执行时间整理成表格，用 Matplotlib 绘制“执行时间—数据规模”双对数图和“比较次数/移动次数”对比图，结合理论复杂度分析实测结果，得出实验结论。"),
]
for head, rest in steps:
    p = c.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.25
    pf.first_line_indent = Pt(BODY * 2)
    p.alignment = JUSTIFY
    r1 = p.add_run(head)
    set_run_font(r1, HEI, EN, BODY, bold=True)
    r2 = p.add_run(rest)
    set_run_font(r2, SONG, EN, BODY)

h2(c, "（二）关键实现说明")
plain(c, "1. 随机数可复现性：srand() 使用固定种子，rand() 在 Windows 下 RAND_MAX = 32767，取值只有 32768 个，映射到 [1, 100000] 上必然出现重复元素，即生成的随机整数允许重复，符合实验要求；固定种子使不同算法面对完全相同的数据。")
plain(c, "2. 统计口径一致性：三种排序算法采用同一套“比较一次记 1 次、写一次数组记 1 次移动”的口径（一次交换记 3 次移动），因此比较次数与移动次数之间具有可比性。")
plain(c, "3. 计时精度与稳定性：采用“同一数组连续排序 repeat 次取平均 + 每种算法测 3 次再取平均”的双重平均策略，既避开了时钟分辨率限制，又削弱了系统调度带来的偶然波动。")
plain(c, "4. 改进蛮力法的正确性保证：先由约束方程消元得到 c = 6a + 3b、7a + 4b = n，再对 a 在 [0, n/7] 内枚举，并用 rest % 4 == 0 判断 b 是否为整数，最后仍保留 b ≥ 0、c ≥ 0、c % 3 == 0 的完整校验，保证与基础蛮力法解集完全一致。")
plain(c, "5. 编码规范：题目1未调用任何标准库排序函数（std::sort、qsort 等）；题目2使用数组存储多组解、未使用全局变量，均满足实验要求。")

h2(c, "（三）完整源代码")
h3(c, "附录 A  实验题目1：冒泡、选择、归并排序算法性能测试程序（exp1_sort.cpp）")
add_code(c, read_text(os.path.join(HERE, "exp1_sort.cpp")))
h3(c, "附录 B  实验题目2：百鸡问题基础蛮力法与改进蛮力法对比程序（exp2_chicken.cpp）")
add_code(c, read_text(os.path.join(HERE, "exp2_chicken.cpp")))

# ----------------------------------------------------------------------------- 实验小结
c = table.rows[8].cells[0]
for p in list(c.paragraphs)[1:]:
    del_par(p)
h2(c, "（一）主要结论")
for t in [
    "1. 三种排序算法与两种百鸡问题求解算法均实现正确：9 组排序测试全部通过“非递减有序”校验，3 组百鸡问题测试中改进蛮力法与基础蛮力法求得的解集逐组一致。",
    "2. 时间复杂度决定算法性能上限。冒泡、选择排序为 O(n²)，n = 10000 时耗时约 6.2×10⁴ 微秒和 9.2×10⁴ 微秒；归并排序为 O(n log n)，仅需 333 微秒，比冒泡排序快约 186.2 倍、比选择排序快约 276.5 倍。",
    "3. 百鸡问题中，通过约束消元把双重循环降为单重循环、把枚举范围压缩到 [0, n/7]，时间复杂度由 O(n²) 降为 O(n)，n = 1000 时循环次数降低 7007 倍、执行时间降低 5471 倍。",
    "4. 优化效果随问题规模增大而增强，说明降低时间复杂度和缩小搜索空间是提升程序效率最有效的手段；改进版冒泡排序的提前终止属于常数级优化，对随机大规模数据收益有限。",
]:
    plain(c, t)
h2(c, "（二）收获与体会")
for t in [
    "通过本次实验，把课堂上“时间复杂度”的抽象概念落实到了可测量的数据上：同一条 n(n−1)/2 的比较次数公式，在程序里是可以被精确统计并与实测耗时相互印证的；而归并排序耗时随规模近似线性对数增长、两种蛮力排序近似平方增长的差异，在双对数坐标图上一目了然。",
    "在实现过程中也体会到“测量方法”本身的重要性：最初直接测量单次排序耗时几乎全部为 0 微秒，改用重复执行取平均后才得到稳定数据。这说明算法性能分析既要懂算法，也要懂测量。",
]:
    plain(c, t)
h2(c, "（三）存在的问题与改进方向")
for t in [
    "1. 测试数据分布单一：题目1只测试了完全随机的数组，未覆盖正序、逆序、大量重复等情形，而改进版冒泡排序的收益恰恰与数据初始有序程度密切相关，后续可补充不同分布下的对比实验。",
    "2. 随机数质量有限：rand() 取值只有 32768 个，映射到 [1, 100000] 区间后重复元素较多，若需更接近真实随机分布，可改用 C++11 <random> 中的 mt19937。",
    "3. 计时仍受运行环境影响：虽然采用了多次平均，但 n = 10000 时归并排序的耗时波动仍较明显，后续可增加重复次数、报告标准差，或在不同机器上对比验证。",
    "4. 尚未扩展到更优算法：可进一步实现堆排序、快速排序（含三数取中、随机化）以及非递归归并排序，比较不同优化策略的效果。",
]:
    plain(c, t)

doc.save(OUT)
print("saved:", OUT)
