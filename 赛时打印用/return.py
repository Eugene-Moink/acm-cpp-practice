import os
from docx import Document
from docx.shared import Pt

# 自动获取当前脚本所在的目录
current_dir = os.path.dirname(os.path.abspath(__file__))
# 拼接绝对路径（确保文件名和这里的完全一致，包括大小写）
md_file_path = os.path.join(current_dir, '个人模板_打印版.md')

# 检查文件是否存在
if not os.path.exists(md_file_path):
    print(f"❌ 找不到文件：{md_file_path}")
    print("请检查文件名是否正确，或者将 .md 文件移动到当前文件夹！")
    exit()

print("✅ 找到文件，开始转换...")
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Consolas'
style.font.size = Pt(10)

with open(md_file_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_code = False
for line in lines:
    if line.startswith('```'):
        in_code = not in_code
        continue
    if line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif in_code:
        p = doc.add_paragraph(line.rstrip())
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    else:
        doc.add_paragraph(line.rstrip())

doc.save(os.path.join(current_dir, '算法模板.docx'))
print("🎉 转换成功！已生成 算法模板.docx")